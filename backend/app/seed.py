from __future__ import annotations

import io
from pathlib import Path

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.audit import log_event
from app.config import get_settings
from app.config_service import ensure_baseline
from app.ingest import extract_pages
from app.models import (
    Cell,
    ColumnType,
    Document,
    DocumentPage,
    EthicalWall,
    Matter,
    MatterMember,
    ReviewTable,
    Role,
    TableColumn,
    TableRow,
    User,
)
from app.security import hash_password
from app.storage import save_bytes

DEMO_MATTER = "Northwind Procurement v. Contoso Logistics — MSA diligence"
REVIEWER_EMAIL = "reviewer@rockhawk.local"
REVIEWER_PASSWORD = "ReviewerDemo1!"
VIEWER_EMAIL = "viewer@rockhawk.local"
WALLED_EMAIL = "walled@rockhawk.local"

MSA_PAGES = [
    (
        "Contoso Logistics Master Services Agreement",
        "This Master Services Agreement (the \"Agreement\") is entered into as of January 15, 2024 "
        "(the \"Effective Date\"), by and between Northwind Procurement LLC, a Delaware limited "
        "liability company (\"Customer\"), and Contoso Logistics, Inc., a Washington corporation "
        "(\"Provider\").",
    ),
    (
        "Governing Law and Venue",
        "This Agreement shall be governed by the laws of the State of Delaware, without regard to "
        "conflicts of law principles. Exclusive venue lies in the state and federal courts located "
        "in Wilmington, Delaware.",
    ),
    (
        "Term and Termination",
        "Either party may terminate this Agreement for convenience upon thirty (30) days' prior "
        "written notice. Provider may also terminate for material breach if Customer fails to cure "
        "within fifteen (15) days after notice.",
    ),
    (
        "Limitation of Liability",
        "Except for confidentiality breaches or willful misconduct, each party's aggregate liability "
        "under this Agreement shall not exceed $500,000. Neither party is liable for indirect or "
        "consequential damages.",
    ),
    (
        "Confidentiality",
        "Each party shall protect Confidential Information of the other using reasonable care. "
        "Confidentiality obligations survive for three (3) years after termination of this Agreement.",
    ),
]

NDA_TEXT = """NORTHWIND / CONTOSO MUTUAL NONDISCLOSURE AGREEMENT

This Mutual Nondisclosure Agreement is made as of March 3, 2023, by and between
Northwind Procurement LLC and Contoso Logistics, Inc.

1. Purpose. The parties wish to explore a possible logistics services relationship.

2. Confidentiality. Confidential Information means non-public business information
disclosed by either party. Confidentiality obligations remain in effect for two (2) years
after the date of disclosure.

3. Governing Law. This agreement is governed by the laws of the State of Washington.

4. No termination for convenience. This NDA may not be terminated for convenience
while discussions are ongoing; it expires automatically on March 3, 2025.

This document is fictional sample text for the RockHawk demonstration matter.
"""

SOW_TEXT = """STATEMENT OF WORK NO. 1
Northwind Warehouse Overflow Support

This Statement of Work is issued under the Contoso Logistics Master Services Agreement
dated January 15, 2024.

Services. Provider will supply overflow warehouse labor in Tacoma, Washington.

Fees. Customer shall pay a monthly service fee of $42,000.

This SOW does not restate governing law, liability cap, or termination for convenience.
Those terms remain as set forth in the MSA. This SOW is fictional demonstration content.
"""

AMENDMENT_TEXT = """AMENDMENT NO. 1 TO MASTER SERVICES AGREEMENT

The parties amend the January 15, 2024 Master Services Agreement as follows:

Liability cap. The aggregate liability cap in the Limitation of Liability section is
increased from $500,000 to $750,000.

No other terms are modified. This amendment is fictional demonstration content and
does not identify any real client.
"""


def seed_if_needed(db: Session) -> None:
    settings = get_settings()
    admin = db.scalar(select(User).where(User.email == settings.admin_email))
    if admin is None:
        admin = User(
            email=settings.admin_email,
            name="RockHawk Admin",
            hashed_password=hash_password(settings.admin_password),
            role=Role.ADMIN.value,
            must_change_password=True,
        )
        admin.admin_scopes = [
            "config.read",
            "config.write",
            "config.apply",
            "ai.manage",
            "users.manage",
            "audit.read",
        ]
        db.add(admin)
        db.flush()
        log_event(db, action="user.seeded", entity_type="user", entity_id=admin.id, actor_id=admin.id)

    ensure_baseline(db, admin.id)

    if not settings.seed_demo:
        db.commit()
        return

    if db.scalar(select(Matter).where(Matter.name == DEMO_MATTER)):
        db.commit()
        return

    reviewer = _user(db, REVIEWER_EMAIL, "Avery Reviewer", REVIEWER_PASSWORD, Role.REVIEWER.value)
    viewer = _user(db, VIEWER_EMAIL, "Casey Viewer", "ViewerDemo1!", Role.VIEWER.value)
    walled = _user(db, WALLED_EMAIL, "Jordan Walled", "WalledDemo1!", Role.ATTORNEY.value)

    matter = Matter(
        name=DEMO_MATTER,
        description=(
            "Fictional diligence file for a logistics MSA. No real client, matter, or "
            "privileged communication is represented."
        ),
        client_name="Northwind Procurement LLC (fictional)",
        opposing_party="Contoso Logistics, Inc. (fictional)",
        created_by_id=admin.id,
    )
    db.add(matter)
    db.flush()

    for user, role in ((admin, Role.ADMIN.value), (reviewer, Role.REVIEWER.value), (viewer, Role.VIEWER.value)):
        db.add(MatterMember(matter_id=matter.id, user_id=user.id, member_role=role))
    db.add(EthicalWall(matter_id=matter.id, user_id=walled.id, reason="Prior representation of Contoso (demo wall)"))

    documents = [
        _add_pdf(db, matter, admin, "Contoso_MSA.pdf", MSA_PAGES),
        _add_text(db, matter, admin, "Mutual_NDA.txt", NDA_TEXT),
        _add_docx(db, matter, admin, "SOW_1_Overflow.docx", SOW_TEXT),
        _add_text(db, matter, admin, "MSA_Amendment_1.txt", AMENDMENT_TEXT),
    ]

    table = ReviewTable(
        matter_id=matter.id,
        name="MSA diligence grid",
        description="Typed extraction columns for the fictional Northwind/Contoso file.",
        created_by_id=admin.id,
    )
    db.add(table)
    db.flush()

    columns = _default_columns(table.id)
    db.add_all(columns)
    db.flush()
    convenience = next(col for col in columns if col.name == "Termination for convenience")
    notice = next(col for col in columns if col.name == "Termination notice period")
    notice.condition_column_id = convenience.id
    notice.condition_equals = "true"

    for document in documents:
        row = TableRow(table_id=table.id, document_id=document.id, is_hot=document.filename.startswith("Contoso"))
        db.add(row)
        db.flush()
        for column in columns:
            db.add(Cell(row_id=row.id, column_id=column.id))

    log_event(
        db,
        action="demo.seeded",
        entity_type="matter",
        entity_id=matter.id,
        actor_id=admin.id,
        matter_id=matter.id,
        payload={"documents": [d.filename for d in documents]},
    )
    db.commit()


def default_column_specs() -> list[dict]:
    return [
        {
            "name": "Governing law",
            "value_type": ColumnType.TEXT.value,
            "instruction": "Extract the governing-law jurisdiction. Use only the document text. If absent, Not found.",
            "citation_policy": "when_quoting",
            "model_role": "extraction",
            "overwrite_policy": "skip_verified",
            "sort_order": 1,
        },
        {
            "name": "Effective date",
            "value_type": ColumnType.DATE.value,
            "instruction": "Extract the effective or as-of date. Do not invent a date.",
            "sort_order": 2,
        },
        {
            "name": "Termination for convenience",
            "value_type": ColumnType.BOOLEAN.value,
            "instruction": "True only if the document expressly allows termination for convenience.",
            "sort_order": 3,
        },
        {
            "name": "Termination notice period",
            "value_type": ColumnType.TEXT.value,
            "instruction": "If termination for convenience exists, extract the required notice period.",
            "sort_order": 4,
        },
        {
            "name": "Liability cap",
            "value_type": ColumnType.MONEY.value,
            "instruction": "Extract the aggregate contractual liability cap amount if stated.",
            "sort_order": 5,
        },
        {
            "name": "Confidentiality duration",
            "value_type": ColumnType.TEXT.value,
            "instruction": "Extract how long confidentiality obligations survive, if stated.",
            "sort_order": 6,
        },
        {
            "name": "Needs attorney review",
            "value_type": ColumnType.ENUM.value,
            "enum_options": ["Yes", "No"],
            "instruction": "Yes if the document addresses liability, termination, privilege, or personal data; else No. Still not a legal conclusion.",
            "sort_order": 7,
        },
    ]


def _default_columns(table_id: str) -> list[TableColumn]:
    return [TableColumn(table_id=table_id, **spec) for spec in default_column_specs()]


def _user(db: Session, email: str, name: str, password: str, role: str) -> User:
    existing = db.scalar(select(User).where(User.email == email))
    if existing:
        return existing
    user = User(email=email, name=name, hashed_password=hash_password(password), role=role)
    db.add(user)
    db.flush()
    return user


def _persist_document(db: Session, matter: Matter, admin: User, filename: str, content_type: str, data: bytes) -> Document:
    import hashlib

    relative = f"{matter.id}/{filename}"
    save_bytes(relative, data)
    pages = extract_pages(filename, data)
    document = Document(
        matter_id=matter.id,
        filename=filename,
        content_type=content_type,
        storage_path=relative,
        content_hash=hashlib.sha256(data).hexdigest(),
        byte_size=len(data),
        page_count=len(pages),
        uploaded_by_id=admin.id,
    )
    db.add(document)
    db.flush()
    for number, text in pages:
        db.add(DocumentPage(document_id=document.id, page_number=number, text=text, embedding=_toy_embedding(text)))
    return document


def _add_text(db: Session, matter: Matter, admin: User, filename: str, text: str) -> Document:
    return _persist_document(db, matter, admin, filename, "text/plain", text.encode("utf-8"))


def _add_docx(db: Session, matter: Matter, admin: User, filename: str, text: str) -> Document:
    buffer = io.BytesIO()
    doc = DocxDocument()
    for block in text.strip().split("\n\n"):
        doc.add_paragraph(block.strip())
    doc.save(buffer)
    return _persist_document(
        db, matter, admin, filename,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        buffer.getvalue(),
    )


def _add_pdf(db: Session, matter: Matter, admin: User, filename: str, sections: list[tuple[str, str]]) -> Document:
    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    story = []
    pdf = SimpleDocTemplate(buffer, pagesize=letter)
    for title, body in sections:
        story.append(Paragraph(title, styles["Heading1"]))
        story.append(Spacer(1, 12))
        story.append(Paragraph(body, styles["BodyText"]))
        story.append(Spacer(1, 18))
    pdf.build(story)
    return _persist_document(db, matter, admin, filename, "application/pdf", buffer.getvalue())


def _toy_embedding(text: str) -> list[float]:
    values = [0.0] * 32
    for index, char in enumerate(text.encode("utf-8")[:256]):
        values[index % 32] += (char % 13) / 13.0
    return values


def write_demo_fixtures(target: Path) -> None:
    target.mkdir(parents=True, exist_ok=True)
    (target / "Mutual_NDA.txt").write_text(NDA_TEXT)
    (target / "MSA_Amendment_1.txt").write_text(AMENDMENT_TEXT)
    (target / "SOW_1_Overflow.txt").write_text(SOW_TEXT)
