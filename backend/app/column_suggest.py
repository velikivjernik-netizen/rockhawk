"""Propose review columns from natural language or checklists. Never persist."""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from openpyxl import load_workbook

VALUE_TYPES = {"text", "boolean", "date", "number", "money", "enum"}

HINTS: list[tuple[tuple[str, ...], dict[str, Any]]] = [
    (("governing law", "choice of law", "jurisdiction"), {
        "name": "Governing law",
        "value_type": "text",
        "instruction": "Extract the governing-law jurisdiction. Use only the document text. If absent, Not found.",
    }),
    (("venue", "forum"), {
        "name": "Venue",
        "value_type": "text",
        "instruction": "Extract the exclusive venue or forum if stated. If absent, Not found.",
    }),
    (("effective date", "as of date", "dated"), {
        "name": "Effective date",
        "value_type": "date",
        "instruction": "Extract the effective or as-of date. Do not invent a date.",
    }),
    (("termination for convenience", "convenience termination"), {
        "name": "Termination for convenience",
        "value_type": "boolean",
        "instruction": "True only if the document expressly allows termination for convenience.",
    }),
    (("notice period", "prior written notice"), {
        "name": "Termination notice period",
        "value_type": "text",
        "instruction": "If termination for convenience exists, extract the required notice period.",
    }),
    (("liability cap", "limitation of liability", "aggregate liability"), {
        "name": "Liability cap",
        "value_type": "money",
        "instruction": "Extract the aggregate contractual liability cap amount if stated.",
    }),
    (("confidential", "nondisclosure"), {
        "name": "Confidentiality duration",
        "value_type": "text",
        "instruction": "Extract how long confidentiality obligations survive, if stated.",
    }),
    (("jury waiver", "waive jury"), {
        "name": "Jury waiver",
        "value_type": "boolean",
        "instruction": "True only if the document expressly waives jury trial.",
    }),
    (("assignment", "assign"), {
        "name": "Assignment restriction",
        "value_type": "text",
        "instruction": "Extract any restriction on assignment. If absent, Not found.",
    }),
    (("indemnif",), {
        "name": "Indemnity",
        "value_type": "text",
        "instruction": "Summarize the indemnity obligation in one sentence from the document only.",
    }),
    (("personal data", "gdpr", "ccpa"), {
        "name": "Privacy / personal data",
        "value_type": "boolean",
        "instruction": "True if the document addresses personal data, GDPR, or CCPA.",
    }),
]


def suggest_from_text(description: str, *, model_role: str = "extraction", citation_policy: str = "when_quoting") -> list[dict[str, Any]]:
    text = description.lower()
    proposed: list[dict[str, Any]] = []
    seen: set[str] = set()
    for needles, spec in HINTS:
        if any(needle in text for needle in needles):
            item = _column(**spec, model_role=model_role, citation_policy=citation_policy)
            if item["name"] not in seen:
                seen.add(item["name"])
                proposed.append(item)
    for chunk in _split_requests(description):
        name = _title(chunk)
        if name and name not in seen and len(name) < 80:
            seen.add(name)
            proposed.append(
                _column(
                    name=name,
                    value_type=_guess_type(chunk),
                    instruction=f"Extract {chunk.strip().rstrip('.')}. Use only the document text. If absent, Not found.",
                    model_role=model_role,
                    citation_policy=citation_policy,
                )
            )
    if not proposed:
        proposed.append(
            _column(
                name=_title(description) or "Custom question",
                value_type="text",
                instruction=description.strip(),
                model_role=model_role,
                citation_policy=citation_policy,
            )
        )
    for index, item in enumerate(proposed, start=1):
        item["sort_order"] = index
    return proposed[:24]


def propose_from_tabular(filename: str, data: bytes) -> list[dict[str, Any]]:
    rows = _read_rows(filename, data)
    if not rows:
        return []
    header = rows[0]
    instruction_row = rows[1] if len(rows) > 1 and _looks_like_instructions(rows[1]) else None
    type_row = None
    for candidate in rows[1:3]:
        if candidate and str(candidate[0]).strip().lower() in VALUE_TYPES:
            type_row = candidate
            break
    proposed: list[dict[str, Any]] = []
    for index, raw_name in enumerate(header):
        name = str(raw_name or "").strip()
        if not name or name.lower() in {"document", "filename", "file", "row"}:
            continue
        instruction = ""
        if instruction_row and index < len(instruction_row):
            instruction = str(instruction_row[index] or "").strip()
        value_type = "text"
        if type_row and index < len(type_row) and str(type_row[index]).strip().lower() in VALUE_TYPES:
            value_type = str(type_row[index]).strip().lower()
        else:
            value_type = _guess_type(name + " " + instruction)
        proposed.append(
            _column(
                name=name,
                value_type=value_type,
                instruction=instruction or f"Extract {name}. Use only the document text. If absent, Not found.",
                sort_order=len(proposed) + 1,
            )
        )
    return proposed


def _column(**kwargs: Any) -> dict[str, Any]:
    return {
        "name": kwargs["name"],
        "value_type": kwargs.get("value_type", "text"),
        "instruction": kwargs.get("instruction", ""),
        "enum_options": kwargs.get("enum_options"),
        "condition_column_id": kwargs.get("condition_column_id"),
        "condition_equals": kwargs.get("condition_equals"),
        "sort_order": kwargs.get("sort_order", 0),
        "citation_policy": kwargs.get("citation_policy", "when_quoting"),
        "model_role": kwargs.get("model_role", "extraction"),
        "prompt_key": kwargs.get("prompt_key", "column.extract"),
        "prompt_version": kwargs.get("prompt_version"),
        "overwrite_policy": kwargs.get("overwrite_policy", "skip_verified"),
        "required": kwargs.get("required", False),
        "validation_json": kwargs.get("validation_json"),
    }


def _split_requests(description: str) -> list[str]:
    parts = re.split(r"[;\n]| and |, ", description)
    return [part.strip(" .") for part in parts if len(part.strip()) > 8]


def _title(text: str) -> str:
    cleaned = re.sub(r"^(extract|find|show|is there|whether)\s+", "", text.strip(), flags=re.I)
    cleaned = cleaned.split(".")[0].strip()
    return cleaned[:80].title() if cleaned else ""


def _guess_type(text: str) -> str:
    lower = text.lower()
    if any(word in lower for word in ("true", "false", "whether", "is there", "yes/no", "boolean")):
        return "boolean"
    if "date" in lower:
        return "date"
    if any(word in lower for word in ("cap", "amount", "fee", "dollar", "money", "$")):
        return "money"
    if "enum" in lower or "yes or no" in lower:
        return "enum"
    return "text"


def _looks_like_instructions(row: list[Any]) -> bool:
    joined = " ".join(str(cell or "") for cell in row).lower()
    return any(word in joined for word in ("extract", "if absent", "not found", "true only", "instruction"))


def _read_rows(filename: str, data: bytes) -> list[list[Any]]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        return [[cell for cell in row] for row in sheet.iter_rows(values_only=True) if any(row)]
    if name.endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
        return [sheet.row_values(index) for index in range(sheet.nrows)]
    text = data.decode("utf-8-sig", errors="replace")
    return list(csv.reader(io.StringIO(text)))
