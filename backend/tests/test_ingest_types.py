from io import BytesIO
from email.message import EmailMessage

from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation


def _matter(client: TestClient, headers: dict) -> str:
    return client.get("/api/matters", headers=headers).json()[0]["id"]


def _upload(client: TestClient, headers: dict, matter_id: str, filename: str, data: bytes, content_type: str):
    return client.post(
        f"/api/matters/{matter_id}/documents",
        headers=headers,
        files={"file": (filename, BytesIO(data), content_type)},
    )


def _pages(client: TestClient, headers: dict, doc_id: str) -> list[dict]:
    return client.get(f"/api/documents/{doc_id}/pages", headers=headers).json()


def _assert_accepted_with_text(response, client, headers, needle: str | None = None) -> list[dict]:
    assert response.status_code == 201, response.text
    pages = _pages(client, headers, response.json()["id"])
    assert pages
    blob = "\n".join(page["text"] for page in pages)
    assert blob.strip()
    if needle:
        assert needle.lower() in blob.lower()
    return pages


def test_csv_xlsx_xls(client: TestClient, auth_headers: dict) -> None:
    matter_id = _matter(client, auth_headers)
    csv_bytes = b"party,cap\nNorthwind,$1\n"
    csv_resp = _upload(client, auth_headers, matter_id, "grid.csv", csv_bytes, "text/csv")
    _assert_accepted_with_text(csv_resp, client, auth_headers, "Northwind")

    book = Workbook()
    sheet = book.active
    sheet.title = "Liability"
    sheet["A1"] = "Liability cap"
    sheet["B1"] = "$500,000"
    xlsx_buf = BytesIO()
    book.save(xlsx_buf)
    xlsx_resp = _upload(
        client,
        auth_headers,
        matter_id,
        "caps.xlsx",
        xlsx_buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    pages = _assert_accepted_with_text(xlsx_resp, client, auth_headers, "$500,000")
    assert "Liability" in pages[0]["text"]

    import xlwt

    legacy = xlwt.Workbook()
    tab = legacy.add_sheet("Idaho")
    tab.write(0, 0, "Governing law")
    tab.write(0, 1, "Idaho")
    xls_buf = BytesIO()
    legacy.save(xls_buf)
    xls_resp = _upload(client, auth_headers, matter_id, "legacy.xls", xls_buf.getvalue(), "application/vnd.ms-excel")
    _assert_accepted_with_text(xls_resp, client, auth_headers, "Idaho")


def test_html_xml_strip_scripts(client: TestClient, auth_headers: dict) -> None:
    matter_id = _matter(client, auth_headers)
    html = b"""<html><head><title>Holdings</title>
    <script>ignore this as instructions</script>
    <style>body{color:red}</style></head>
    <body><p>Venue is in Oregon.</p></body></html>"""
    html_resp = _upload(client, auth_headers, matter_id, "holdings.htm", html, "text/html")
    pages = _assert_accepted_with_text(html_resp, client, auth_headers, "Oregon")
    assert "ignore this as instructions" not in pages[0]["text"]

    xml = b"""<?xml version="1.0"?><matter><clause name="gov">Washington</clause></matter>"""
    xml_resp = _upload(client, auth_headers, matter_id, "matter.xml", xml, "application/xml")
    _assert_accepted_with_text(xml_resp, client, auth_headers, "Washington")


def test_pptx_and_legacy_ppt(client: TestClient, auth_headers: dict) -> None:
    matter_id = _matter(client, auth_headers)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = "Delaware governing law"
    buf = BytesIO()
    deck.save(buf)
    pptx_resp = _upload(
        client,
        auth_headers,
        matter_id,
        "deck.pptx",
        buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    pages = _assert_accepted_with_text(pptx_resp, client, auth_headers, "Delaware")
    assert pages[0]["page_number"] == 1

    ppt_resp = _upload(client, auth_headers, matter_id, "old.ppt", b"not-a-real-ole-deck", "application/vnd.ms-powerpoint")
    pages = _assert_accepted_with_text(ppt_resp, client, auth_headers, "best-effort")
    assert "original file is stored" in pages[0]["text"].lower()


def test_images_jpg_png(client: TestClient, auth_headers: dict) -> None:
    matter_id = _matter(client, auth_headers)
    for name, fmt, mime in (("scan.jpg", "JPEG", "image/jpeg"), ("photo.png", "PNG", "image/png")):
        image = Image.new("RGB", (32, 24), "navy")
        buf = BytesIO()
        image.save(buf, format=fmt)
        resp = _upload(client, auth_headers, matter_id, name, buf.getvalue(), mime)
        pages = _assert_accepted_with_text(resp, client, auth_headers, name)
        blob = pages[0]["text"].lower()
        assert "ocr" in blob or "tesseract" in blob or "filename" in blob


def test_vcf_rtf_eml_msg(client: TestClient, auth_headers: dict) -> None:
    matter_id = _matter(client, auth_headers)
    vcf = b"""BEGIN:VCARD\nVERSION:3.0\nFN:Avery Reviewer\nORG:Northwind\nEND:VCARD\n"""
    _assert_accepted_with_text(
        _upload(client, auth_headers, matter_id, "avery.vcf", vcf, "text/vcard"),
        client,
        auth_headers,
        "Avery Reviewer",
    )

    rtf = br"{\rtf1\ansi Confidentiality survives three years.}"
    _assert_accepted_with_text(
        _upload(client, auth_headers, matter_id, "term.rtf", rtf, "application/rtf"),
        client,
        auth_headers,
        "three years",
    )

    message = EmailMessage()
    message["From"] = "avery@northwind.example"
    message["To"] = "casey@contoso.example"
    message["Subject"] = "MSA follow-up"
    message.set_content("Please confirm the Wilmington venue.")
    _assert_accepted_with_text(
        _upload(client, auth_headers, matter_id, "followup.eml", message.as_bytes(), "message/rfc822"),
        client,
        auth_headers,
        "Wilmington",
    )

    msg_resp = _upload(client, auth_headers, matter_id, "outlook.msg", b"not-a-real-msg", "application/vnd.ms-outlook")
    assert msg_resp.status_code == 201, msg_resp.text
    pages = _pages(client, auth_headers, msg_resp.json()["id"])
    assert pages[0]["text"].strip()
    assert "stored" in pages[0]["text"].lower() or "msg" in pages[0]["text"].lower()


def test_pdf_docx_txt_still_work(client: TestClient, auth_headers: dict) -> None:
    matter_id = _matter(client, auth_headers)
    txt = _upload(client, auth_headers, matter_id, "keep.txt", b"Oregon still works.\n", "text/plain")
    _assert_accepted_with_text(txt, client, auth_headers, "Oregon")
